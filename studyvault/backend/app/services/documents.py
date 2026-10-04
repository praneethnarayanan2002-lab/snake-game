"""Document formats: validation by file signature and plain-text extraction.

Only an allowlist of document/image formats is accepted. Formats a browser could
execute (HTML, SVG) and binaries are rejected because uploads are served publicly.
"""

import csv
import io
import logging
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO

from app.services.pdf_text import MAX_TEXT_CHARS, InvalidPdf, extract_pdf

log = logging.getLogger(__name__)


@dataclass(frozen=True)
class Format:
    kind: str  # pdf | document | presentation | spreadsheet | text | image
    label: str
    mime: str


FORMATS: dict[str, Format] = {
    "pdf": Format("pdf", "PDF", "application/pdf"),
    "docx": Format("document", "Word", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "doc": Format("document", "Word (legacy)", "application/msword"),
    "odt": Format("document", "OpenDocument Text", "application/vnd.oasis.opendocument.text"),
    "rtf": Format("document", "Rich Text", "application/rtf"),
    "pptx": Format("presentation", "PowerPoint", "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    "ppt": Format("presentation", "PowerPoint (legacy)", "application/vnd.ms-powerpoint"),
    "odp": Format("presentation", "OpenDocument Presentation", "application/vnd.oasis.opendocument.presentation"),
    "xlsx": Format("spreadsheet", "Excel", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "xls": Format("spreadsheet", "Excel (legacy)", "application/vnd.ms-excel"),
    "ods": Format("spreadsheet", "OpenDocument Spreadsheet", "application/vnd.oasis.opendocument.spreadsheet"),
    "csv": Format("spreadsheet", "CSV", "text/csv"),
    "txt": Format("text", "Text", "text/plain"),
    "md": Format("text", "Markdown", "text/markdown"),
    "png": Format("image", "PNG image", "image/png"),
    "jpg": Format("image", "JPEG image", "image/jpeg"),
    "jpeg": Format("image", "JPEG image", "image/jpeg"),
    "webp": Format("image", "WebP image", "image/webp"),
    "gif": Format("image", "GIF image", "image/gif"),
}

ALLOWED_MIME_TYPES = sorted({f.mime for f in FORMATS.values()})

_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_ZIP_MARKERS = {"docx": "word/document.xml", "pptx": "ppt/presentation.xml", "xlsx": "xl/workbook.xml"}
_ODF_MIMES = {"odt": "application/vnd.oasis.opendocument.text", "odp": "application/vnd.oasis.opendocument.presentation", "ods": "application/vnd.oasis.opendocument.spreadsheet"}


class InvalidDocument(ValueError):
    pass


@dataclass
class DocInfo:
    ext: str
    kind: str
    mime: str
    page_count: int  # pages / slides / sheets; 0 when not meaningful
    text: str


def extension(filename: str) -> str:
    return Path(filename or "").suffix.lower().lstrip(".")


def format_for(filename: str) -> Format:
    ext = extension(filename)
    if ext not in FORMATS:
        supported = ", ".join(sorted(FORMATS))
        raise InvalidDocument(f"Unsupported file type{f' .{ext}' if ext else ''}. Supported: {supported}")
    return FORMATS[ext]


def _clean(text: str) -> str:
    text = re.sub(r"[ \t\x00]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()[:MAX_TEXT_CHARS]


def _check_signature(ext: str, data: bytes) -> None:
    head = data[:16]
    ok = True
    if ext == "pdf":
        ok = head.startswith(b"%PDF-")
    elif ext in ("doc", "ppt", "xls"):
        ok = head.startswith(_OLE)
    elif ext == "rtf":
        ok = head.startswith(b"{\\rtf")
    elif ext == "png":
        ok = head.startswith(b"\x89PNG\r\n\x1a\n")
    elif ext in ("jpg", "jpeg"):
        ok = head.startswith(b"\xff\xd8\xff")
    elif ext == "gif":
        ok = head[:6] in (b"GIF87a", b"GIF89a")
    elif ext == "webp":
        ok = head[:4] == b"RIFF" and head[8:12] == b"WEBP"
    elif ext in _ZIP_MARKERS or ext in _ODF_MIMES:
        ok = False
        if head.startswith(b"PK\x03\x04"):
            try:
                with zipfile.ZipFile(io.BytesIO(data)) as z:
                    names = set(z.namelist())
                    if ext in _ZIP_MARKERS:
                        ok = _ZIP_MARKERS[ext] in names
                    else:
                        ok = "mimetype" in names and z.read("mimetype").decode(errors="ignore").strip() == _ODF_MIMES[ext]
            except zipfile.BadZipFile:
                ok = False
    elif ext in ("txt", "md", "csv"):
        ok = b"\x00" not in data[:8192]
    if not ok:
        raise InvalidDocument(f"This file isn't a valid .{ext} — it may be corrupted or renamed")


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _docx(data: bytes) -> tuple[str, int]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts), 0


def _pptx(data: bytes) -> tuple[str, int]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"— Slide {i} —")
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(parts), len(prs.slides)


def _xlsx(data: bytes) -> tuple[str, int]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts, total = [], 0
    for ws in wb.worksheets:
        parts.append(f"— {ws.title} —")
        for row in ws.iter_rows(values_only=True):
            line = " | ".join("" if v is None else str(v) for v in row).strip(" |")
            if line:
                parts.append(line)
                total += len(line)
            if total > MAX_TEXT_CHARS:
                break
    return "\n".join(parts), len(wb.worksheets)


def _odf(data: bytes) -> tuple[str, int]:
    from odf import teletype
    from odf.opendocument import load
    from odf.text import P

    doc = load(io.BytesIO(data))
    return "\n".join(teletype.extractText(p) for p in doc.getElementsByType(P)), 0


def _rtf(data: bytes) -> tuple[str, int]:
    from striprtf.striprtf import rtf_to_text

    return rtf_to_text(_decode(data)), 0


def _csv(data: bytes) -> tuple[str, int]:
    rows = csv.reader(io.StringIO(_decode(data)))
    return "\n".join(" | ".join(r) for r in rows), 0


_EXTRACTORS = {
    "docx": _docx,
    "pptx": _pptx,
    "xlsx": _xlsx,
    "odt": _odf,
    "odp": _odf,
    "ods": _odf,
    "rtf": _rtf,
    "csv": _csv,
    "txt": lambda d: (_decode(d), 0),
    "md": lambda d: (_decode(d), 0),
}


def extract_document(fileobj: BinaryIO, filename: str) -> DocInfo:
    """Validate an upload and extract its searchable text."""
    fmt = format_for(filename)
    ext = extension(filename)
    fileobj.seek(0)
    if ext == "pdf":
        try:
            info = extract_pdf(fileobj)
        except InvalidPdf as exc:
            raise InvalidDocument(str(exc)) from exc
        return DocInfo(ext, fmt.kind, fmt.mime, info.page_count, info.text)

    data = fileobj.read()
    fileobj.seek(0)
    _check_signature(ext, data)
    text, pages = "", 0
    if extractor := _EXTRACTORS.get(ext):
        try:
            text, pages = extractor(data)
        except Exception as exc:  # noqa: BLE001 - parsers raise many error types on bad files
            log.warning("Extraction failed for .%s", ext, exc_info=True)
            raise InvalidDocument(f"Couldn't read this .{ext} file — it may be corrupted") from exc
    # Legacy binary Office files and images are stored without extracted text.
    return DocInfo(ext, fmt.kind, fmt.mime, pages, _clean(text))
