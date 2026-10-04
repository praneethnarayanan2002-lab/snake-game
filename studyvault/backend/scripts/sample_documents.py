"""Non-PDF seed files (Word, PowerPoint, Excel, image, Markdown, CSV)."""

import io

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches
from pptx.util import Pt as PPt


def sql_joins_docx() -> bytes:
    doc = Document()
    doc.add_heading("SQL Joins — Cheat Sheet", 0)
    doc.add_paragraph("Quick reference for DBMS Unit 2: how each join combines rows from two relations.")
    joins = [
        ("INNER JOIN", "Rows with matching keys in both tables.", "SELECT * FROM student s JOIN enrol e ON s.id = e.sid;"),
        ("LEFT OUTER JOIN", "All rows from the left table, NULLs where the right side has no match.", "SELECT * FROM student s LEFT JOIN enrol e ON s.id = e.sid;"),
        ("RIGHT OUTER JOIN", "All rows from the right table, NULLs where the left side has no match.", "SELECT * FROM student s RIGHT JOIN enrol e ON s.id = e.sid;"),
        ("FULL OUTER JOIN", "Union of left and right outer joins.", "SELECT * FROM a FULL JOIN b ON a.k = b.k;"),
        ("NATURAL JOIN", "Equi-join on all columns with the same name; duplicate columns removed.", "SELECT * FROM student NATURAL JOIN enrol;"),
        ("SELF JOIN", "A table joined with itself, e.g. employee and manager.", "SELECT e.name, m.name FROM emp e JOIN emp m ON e.mgr = m.id;"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, ["Join", "Meaning", "Example"]):
        cell.text = text
    for name, meaning, example in joins:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = name, meaning, example
    doc.add_heading("Exam tips", level=1)
    for tip in ["Draw the Venn diagram for outer joins.", "State that natural join can silently join on unintended columns.", "Write the relational-algebra form: R ⋈ S."]:
        p = doc.add_paragraph(tip, style="List Bullet")
        p.runs[0].font.size = Pt(11)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def cpu_scheduling_pptx() -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
    slides = [
        ("CPU Scheduling", ["Operating Systems · Unit 2", "Lecture slides"]),
        ("Scheduling criteria", ["CPU utilisation and throughput", "Turnaround, waiting and response time", "Fairness and starvation"]),
        ("FCFS", ["Non-preemptive, served in arrival order", "Simple but suffers the convoy effect"]),
        ("SJF / SRTF", ["Shortest job first minimises average waiting time", "SRTF is the preemptive version", "Needs burst-time prediction (exponential averaging)"]),
        ("Round Robin", ["Each process gets a time quantum", "Small quantum → many context switches", "Large quantum → behaves like FCFS"]),
        ("Priority scheduling", ["Low-priority processes may starve", "Aging gradually raises priority"]),
    ]
    for i, (title, bullets) in enumerate(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[0 if i == 0 else 1])
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for b in bullets[1:]:
            body.add_paragraph().text = b
        for p in body.paragraphs:
            for r in p.runs:
                r.font.size = PPt(26)
        slide.notes_slide.notes_text_frame.text = f"Speaker notes: explain {title.lower()} with a Gantt chart."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def distributions_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Distributions"
    rows = [
        ("Distribution", "Parameters", "Mean", "Variance", "Use when"),
        ("Binomial", "n, p", "np", "np(1-p)", "Number of successes in n independent trials"),
        ("Poisson", "λ", "λ", "λ", "Count of rare events in a fixed interval"),
        ("Geometric", "p", "1/p", "(1-p)/p²", "Trials until the first success"),
        ("Uniform (continuous)", "a, b", "(a+b)/2", "(b-a)²/12", "Every value in [a, b] equally likely"),
        ("Exponential", "λ", "1/λ", "1/λ²", "Waiting time between Poisson events"),
        ("Normal", "μ, σ²", "μ", "σ²", "Sums of many small effects; z = (x-μ)/σ"),
    ]
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="5B57E8")
    for col, width in zip("ABCDE", (22, 12, 10, 12, 48)):
        ws.column_dimensions[col].width = width
    z = wb.create_sheet("Z-table (excerpt)")
    z.append(("z", "P(Z ≤ z)"))
    for zval, p in [(0.0, 0.5), (0.5, 0.6915), (1.0, 0.8413), (1.5, 0.9332), (1.96, 0.975), (2.0, 0.9772), (2.5, 0.9938), (3.0, 0.9987)]:
        z.append((zval, p))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def tree_traversal_png() -> bytes:
    img = Image.new("RGB", (1200, 800), "#fbfbf8")
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 34)
        small = ImageFont.truetype("DejaVuSans.ttf", 26)
    except OSError:
        font = small = ImageFont.load_default()
    d.text((60, 40), "Binary tree traversals", fill="#1f2937", font=font)
    nodes = {"A": (600, 170), "B": (380, 330), "C": (820, 330), "D": (260, 500), "E": (500, 500), "F": (940, 500)}
    for a, b in [("A", "B"), ("A", "C"), ("B", "D"), ("B", "E"), ("C", "F")]:
        d.line([nodes[a], nodes[b]], fill="#6b7280", width=4)
    for name, (x, y) in nodes.items():
        d.ellipse([x - 42, y - 42, x + 42, y + 42], fill="#ffffff", outline="#5b57e8", width=5)
        d.text((x - 11, y - 18), name, fill="#111827", font=font)
    for i, line in enumerate(["Inorder:   D B E A C F", "Preorder:  A B D E C F", "Postorder: D E B F C A", "Level:     A B C D E F"]):
        d.text((60, 600 + i * 42), line, fill="#374151", font=small)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def subnetting_md() -> bytes:
    return """# Subnetting practice problems

Computer Networks · Unit 3. Solutions at the end.

1. Divide 192.168.10.0/24 into 4 equal subnets. Give the network address, broadcast address and host range of each.
2. How many usable hosts does a /27 subnet have?
3. A company needs 50 hosts per subnet. What is the longest prefix that works?
4. Is 10.0.5.130 in the same subnet as 10.0.5.100 with mask 255.255.255.192?
5. Summarise 172.16.0.0/24 to 172.16.3.0/24 into a single route.

## Solutions

1. /26 subnets: .0, .64, .128, .192 — each with 62 usable hosts.
2. 2^5 − 2 = 30 hosts.
3. /26 (62 hosts).
4. No — .100 is in .64/26 and .130 is in .128/26.
5. 172.16.0.0/22.
""".encode()


def tcp_udp_csv() -> bytes:
    return """Feature,TCP,UDP
Connection,Connection-oriented (3-way handshake),Connectionless
Reliability,Acknowledgements and retransmission,Best effort
Ordering,In-order delivery,No ordering guarantee
Flow control,Sliding window,None
Congestion control,Slow start / AIMD,None
Header size,20–60 bytes,8 bytes
Typical uses,"HTTP, SMTP, FTP","DNS, VoIP, streaming, games"
""".encode()


# subject code, unit, resource type, year, title, description, tags, filename, builder
SAMPLES = [
    ("DBMS", 2, "notes", 2025, "SQL Joins — Cheat Sheet (Word)", "One-page Word cheat sheet of every SQL join with examples and exam tips.", ["dbms", "sql", "joins", "cheat sheet"], "sql-joins-cheat-sheet.docx", sql_joins_docx),
    ("OS", 2, "notes", 2025, "CPU Scheduling — Lecture Slides", "Lecture deck on FCFS, SJF, SRTF, Round Robin and priority scheduling, with speaker notes.", ["os", "cpu scheduling", "round robin", "slides"], "cpu-scheduling-slides.pptx", cpu_scheduling_pptx),
    ("MATH", 3, "reference", 2024, "Probability Distributions — Formula Table", "Excel workbook: mean/variance of common distributions plus a z-table excerpt.", ["maths", "probability", "distributions", "formula sheet"], "probability-distributions.xlsx", distributions_xlsx),
    ("DS", 3, "notes", 2025, "Binary Tree Traversals — Whiteboard Diagram", "Diagram of inorder, preorder, postorder and level-order traversals on one tree.", ["ds", "trees", "traversal", "diagram"], "tree-traversals.png", tree_traversal_png),
    ("CN", 3, "question_bank", 2025, "Subnetting Practice Problems (Markdown)", "Five subnetting problems with worked answers.", ["cn", "subnetting", "ipv4", "practice"], "subnetting-practice.md", subnetting_md),
    ("CN", 4, "reference", 2024, "TCP vs UDP — Comparison Table (CSV)", "Side-by-side comparison of TCP and UDP features.", ["cn", "tcp", "udp", "comparison"], "tcp-vs-udp.csv", tcp_udp_csv),
]
